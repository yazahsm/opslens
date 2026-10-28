"""Excel-compatible CSV and formatted XLSX reports for this application."""
import csv
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from analytics import DEFINITIONS


def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def csv_bytes(rows, headers):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([safe(row.get(key)) for key in headers])
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def rejected_rows(dataset):
    return [{"source_row": row["row"], "record_key": row["key"], "reason": "; ".join(row["reasons"]),
             **row["raw"]} for row in dataset["rejected"]]


def excel_bytes(summary, orders, inventory, datasets):
    book = Workbook()
    book.remove(book.active)
    navy, teal, muted = "172E35", "0E8074", "617578"
    created = datetime.now().astimezone().isoformat(timespec="seconds")

    def sheet(name, title, subtitle, headers, rows, widths=None):
        ws = book.create_sheet(name)
        count = len(headers)
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, count))
        ws.cell(1, 1, title).font = Font(name="Calibri", size=19, bold=True, color=navy)
        ws.row_dimensions[1].height = 32
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, count))
        ws.cell(2, 1, subtitle).font = Font(name="Calibri", size=10, color=muted)
        ws.row_dimensions[2].height = 28
        for column, header in enumerate(headers, 1):
            cell = ws.cell(4, column, header.replace("_", " ").title())
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(column)].width = widths[column-1] if widths else (22 if column <= 4 else 18)
        ws.row_dimensions[4].height = 25
        for index, record in enumerate(rows, 5):
            for column, value in enumerate(record, 1):
                cell = ws.cell(index, column, safe(value))
                cell.font = Font(name="Calibri", size=11, color=navy)
                cell.alignment = Alignment(vertical="top")
                if index % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor="F1F7F5")
                if isinstance(value, float):
                    cell.number_format = "#,##0.00"
        ws.freeze_panes = "A5"
        if rows:
            ws.auto_filter.ref = f"A4:{get_column_letter(count)}{len(rows)+4}"
        return ws

    period = summary["period"]
    subtitle = f"{period['start']} to {period['end']} | {period['region'] or 'All regions'} | {period['category'] or 'All categories'} | INR"
    metric_rows = []
    labels = {"sales": "Net sales (INR)", "orders": "Orders", "on_time": "On-time delivery (%)", "margin": "Gross margin (%)"}
    for key, label in labels.items():
        metric_rows.append([label, summary["kpis"][key], summary["previous"][key], summary["deltas"][key],
                            "percentage points" if key in ("on_time", "margin") else "% change"])
    ws = sheet("Summary", "Operations performance report", subtitle,
               ["Metric", "Current period", "Previous period", "Change", "Change unit"], metric_rows, [28, 24, 24, 18, 24])
    ws.cell(12, 1, "Generated at"); ws.cell(12, 2, created)
    ws.cell(13, 1, "Previous period"); ws.cell(13, 2, f"{period['previous_start']} to {period['previous_end']}")
    ws.cell(15, 1, "Scope").font = Font(bold=True, color=teal)
    ws.cell(16, 1, "Orders and trend use the selected date, region and category filters.")
    ws.cell(17, 1, "Inventory is a separate snapshot; the order date filter does not apply.")
    ws.cell(18, 1, "Data quality and rejected rows describe the full active source datasets.")
    ws.cell(20, 1, "Source types"); ws.cell(20, 2, ", ".join(f"{d['kind']}: {'fictional demo' if d['is_demo'] else 'imported file'}" for d in datasets))
    trend = summary["trend"]
    sheet("Sales trend", "Daily delivered sales", subtitle, ["date", "sales_inr", "previous_period_sales_inr"],
          [[datetime.fromisoformat(row["date"]), row["sales"], row["previous"]] for row in trend], [20, 24, 28])
    for row in book["Sales trend"].iter_rows(min_row=5, max_col=1):
        row[0].number_format = "dd mmm yyyy"
    category_rows = [[row["category"], row["sales"], row["orders"]] for row in summary["categories"]]
    categories = sheet("Categories", "Sales by category", subtitle, ["category", "sales_inr", "delivered_orders"], category_rows, [26, 22, 22])
    if category_rows:
        chart = BarChart()
        chart.title = "Delivered sales by category"
        chart.y_axis.title = "INR"
        chart.add_data(Reference(categories, min_col=2, min_row=4, max_row=4+len(category_rows)), titles_from_data=True)
        chart.set_categories(Reference(categories, min_col=1, min_row=5, max_row=4+len(category_rows)))
        chart.width, chart.height = 20, 10
        categories.add_chart(chart, "E4")
    headers = ["order_id", "order_date", "product", "region", "category", "channel", "units", "unit_price", "unit_cost", "status", "delivery_days", "target_delivery_days", "order_value", "net_sales"]
    sheet("Orders", "Valid orders in selected scope", subtitle, headers, [[row.get(key) for key in headers] for row in orders])
    inv_headers = ["sku", "product", "region", "category", "stock_on_hand", "reorder_point", "avg_daily_units", "coverage_days", "stock_value", "snapshot_date"]
    sheet("Inventory", "Inventory snapshot", "Region and category filters apply. Order dates do not apply.", inv_headers,
          [[row.get(key) for key in inv_headers] for row in inventory["rows"]])
    quality = [[d["kind"], d["name"], d["raw_count"], d["accepted_count"], d["rejected_count"], d["normalized_count"], d["pass_rate"]] for d in datasets]
    sheet("Data quality", "Source validation", "Dataset-wide validation, before dashboard filters. Pass rate = accepted / input rows x 100.",
          ["dataset", "source", "input_rows", "accepted_rows", "excluded_rows", "normalized_rows", "pass_rate_pct"], quality)
    rejected = [[d["kind"], row["row"], row["key"], "; ".join(row["reasons"])] for d in datasets for row in d["rejected"]]
    sheet("Excluded rows", "Excluded source records", "Dataset-wide exclusions. Original values can be downloaded from the dashboard as CSV.",
          ["dataset", "source_row", "record_key", "reason"], rejected, [18, 18, 30, 100])
    sheet("Definitions", "Metric definitions", "All amounts are INR. Source records are one order per product, not order lines.",
          ["metric", "definition"], DEFINITIONS, [28, 115])
    for row in book["Definitions"].iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        book["Definitions"].row_dimensions[row[0].row].height = 42
    sources = [[d["kind"], d["name"], d["id"], d["created_at"], d["source_hash"], "Fictional generated data" if d["is_demo"] else "User-imported local file"] for d in datasets]
    sheet("Sources", "Report provenance", "Dataset IDs and SHA-256 digests identify the validated source revisions used for this report.",
          ["dataset", "source_name", "revision_id", "imported_at", "sha256", "source_type"], sources, [18, 34, 38, 34, 72, 32])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
