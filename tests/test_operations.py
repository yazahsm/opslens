"""Behavior tests for calculations, validation, persistence, and real HTTP routes."""
import csv
from datetime import date, timedelta
import io
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from analytics import inventory_summary, summarize
from app import create_app
from data import clean, columns, parse_file
from demo import generate_demo
from fastapi.testclient import TestClient
from reports import csv_bytes


def order(identifier="A", day=None, **updates):
    row = dict(order_id=identifier, order_date=day or date.today().isoformat(), sku="SKU-1", product="Desk lamp",
               region="South", category="Home", channel="Online", units=2, unit_price=100, unit_cost=60,
               status="Delivered", delivery_days=2, target_delivery_days=3)
    return {**row, **updates}


class CalculationTests(unittest.TestCase):
    def test_kpis_are_reconciled_and_rates_use_correct_denominators(self):
        today = date.today()
        yesterday = (today-timedelta(days=1)).isoformat()
        rows = [order("prior", yesterday, units=1), order("one"), order("two", units=1, delivery_days=5),
                order("cancel", units=10, status="Cancelled", delivery_days=""),
                order("processing", units=10, status="Processing", delivery_days="")]
        data = clean(rows, "orders")
        summary = summarize(data, start=today.isoformat(), end=today.isoformat())
        self.assertEqual(summary["kpis"]["sales"], 300)
        self.assertEqual(summary["kpis"]["orders"], 4)
        self.assertEqual(summary["kpis"]["on_time"], 50)
        self.assertEqual(summary["kpis"]["margin"], 40)
        self.assertEqual(summary["kpis"]["gross_profit"], 120)
        self.assertEqual(summary["deltas"]["sales"], 200)
        self.assertEqual(summary["deltas"]["on_time"], -50)
        self.assertEqual(summary["trend"][0]["previous"], 100)
        self.assertEqual(summary["categories"][0]["sales"], 300)

    def test_empty_selection_and_zero_sales_never_produce_nan(self):
        data = clean([order(unit_price=0)], "orders")
        summary = summarize(data)
        self.assertIsNone(summary["kpis"]["margin"])
        self.assertIsNone(summary["deltas"]["sales"])
        empty = summarize(data, region="No matching region")
        self.assertEqual(empty["kpis"]["orders"], 0)
        self.assertIsNone(empty["kpis"]["on_time"])
        self.assertEqual(empty["categories"], [])
        with self.assertRaises(ValueError):
            summarize(data, start="2026-02-01", end="2026-01-01")

    def test_validation_counts_rows_once_and_retains_reasons(self):
        rows = [order("A", region=" south "), order("A"), order("B", units=-1, order_date="bad-date"),
                order("C", unit_price="NaN"), order("D", delivery_days=""), order("E", units=1.5),
                order("F", unit_price=1.001)]
        result = clean(rows, "orders")
        self.assertEqual((result["accepted_count"], result["rejected_count"], result["normalized_count"]), (1, 6, 1))
        self.assertEqual(result["rows"][0]["region"], "South")
        self.assertEqual(result["rejected"][0]["row"], 3)
        self.assertGreaterEqual(len(result["rejected"][1]["reasons"]), 2)
        self.assertEqual(result["rejected"][1]["raw"]["units"], -1)

    def test_invalid_first_record_does_not_hide_later_valid_record(self):
        result = clean([order("same", units=-1), order("same")], "orders")
        self.assertEqual(result["accepted_count"], 1)
        self.assertEqual(result["rejected_count"], 1)

    def test_inventory_snapshot_totals_and_zero_demand(self):
        _, inventory = generate_demo()
        row = {**inventory[0], "stock_on_hand": 0, "avg_daily_units": 0}
        data = clean([row], "inventory")
        result = inventory_summary(data)
        self.assertEqual(result["low_stock"], 1)
        self.assertEqual(result["stock_value"], 0)
        self.assertIsNone(result["rows"][0]["coverage_days"])

    def test_anomaly_uses_prior_calendar_days(self):
        today = date.today()
        rows = [order(str(i), (today-timedelta(days=7-i)).isoformat(), units=1, unit_price=100 if i<7 else 10) for i in range(8)]
        result = summarize(clean(rows,"orders"), start=today.isoformat(), end=today.isoformat())
        self.assertEqual(result["anomalies"][0]["change"], -90)
        self.assertEqual(result["anomalies"][0]["baseline"], 100)


class APITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database = Path(self.temp.name) / "test.db"
        self.app = create_app(self.database)
        self.client = TestClient(self.app)

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def preview(self, rows, kind="orders"):
        payload=csv_bytes(rows, columns(kind))
        response=self.client.post("/api/imports/preview?kind="+kind, files={"file":("source.csv",payload,"text/csv")})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def test_demo_routes_and_inventory_scope(self):
        self.assertEqual(self.client.get("/").status_code,200)
        self.assertEqual(self.client.get("/api/health").json()["status"],"ok")
        meta=self.client.get("/api/meta").json()
        orders=next(d for d in meta["datasets"] if d["kind"]=="orders")
        self.assertEqual(orders["rejected_count"],27)
        summary=self.client.get("/api/overview?region=South&category=Home").json()
        result=self.client.get("/api/orders?region=South&category=Home&page_size=100").json()
        self.assertEqual(summary["kpis"]["orders"],result["total"])
        self.assertTrue(all(r["region"]=="South" and r["category"]=="Home" for r in result["rows"]))
        inventory1=self.client.get("/api/inventory?days=7").json()
        inventory2=self.client.get("/api/inventory?days=90").json()
        self.assertEqual(inventory1,inventory2)
        quality=self.client.get("/api/quality?region=South").json()
        self.assertEqual(quality["rejected_count"],27)
        self.assertEqual(self.client.get("/api/overview?start=invalid").status_code,400)
        self.assertEqual(self.client.get("/api/quality?kind=unknown").status_code,400)

    def test_import_preview_is_non_mutating_apply_persists_and_restore_works(self):
        original=self.client.get("/api/meta").json()["datasets"][0]["id"]
        preview=self.preview([order("import-one"),order("import-one"),order("invalid",units=-2)])
        self.assertEqual((preview["accepted_count"],preview["rejected_count"]),(1,2))
        self.assertEqual(self.client.get("/api/meta").json()["datasets"][0]["id"],original)
        applied=self.client.post("/api/imports/apply",json={"token":preview["token"]})
        self.assertEqual(applied.status_code,200)
        self.assertEqual(self.client.get("/api/overview").json()["kpis"]["orders"],1)
        self.assertEqual(self.client.post("/api/imports/apply",json={"token":preview["token"]}).status_code,400)
        restarted=TestClient(create_app(self.database))
        self.assertEqual(restarted.get("/api/overview").json()["kpis"]["orders"],1)
        restarted.close()
        self.assertEqual(self.client.post(f"/api/history/{original}/activate").status_code,200)
        self.assertGreater(self.client.get("/api/overview").json()["kpis"]["orders"],1)

    def test_invalid_file_rejected_without_destroying_dataset(self):
        before=self.client.get("/api/meta").json()["datasets"][0]["id"]
        for name,content in [("bad.csv",b"one,two\n1,2"),("bad.exe",b"not a csv"),("bad.csv",b"order_id,order_id\na,b")]:
            response=self.client.post("/api/imports/preview",files={"file":(name,content)})
            self.assertEqual(response.status_code,400,response.text)
        preview=self.preview([order(units=-1)])
        self.assertEqual(self.client.post("/api/imports/apply",json={"token":preview["token"]}).status_code,400)
        self.assertEqual(self.client.get("/api/meta").json()["datasets"][0]["id"],before)

    def test_xlsx_import_dates_and_formula_cells(self):
        book=Workbook();sheet=book.active;sheet.append(columns("orders"))
        row=order("xlsx-row");row["order_date"]=date.today()
        sheet.append([row[key] for key in columns("orders")])
        other=order("formula-row",unit_price="=10*10")
        sheet.append([other[key] for key in columns("orders")])
        output=io.BytesIO();book.save(output)
        response=self.client.post("/api/imports/preview",files={"file":("orders.xlsx",output.getvalue())})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()["accepted_count"],1)
        self.assertEqual(response.json()["rejected_count"],1)

    def test_report_workbook_matches_filtered_kpis_and_is_safe(self):
        response=self.client.get("/api/export/report.xlsx?region=South&category=Home&days=7")
        self.assertEqual(response.status_code,200)
        workbook=load_workbook(io.BytesIO(response.content),data_only=True)
        self.assertEqual(len(workbook.sheetnames),9)
        summary=self.client.get("/api/overview?region=South&category=Home&days=7").json()
        self.assertEqual(workbook["Summary"]["B5"].value,summary["kpis"]["sales"])
        self.assertEqual(workbook["Summary"]["B6"].value,summary["kpis"]["orders"])
        sales=sum(row[1] or 0 for row in workbook["Sales trend"].iter_rows(min_row=5,values_only=True))
        self.assertAlmostEqual(sales,summary["kpis"]["sales"],places=2)
        for row in workbook["Orders"].iter_rows(min_row=5,values_only=True):
            self.assertEqual(row[3],"South");self.assertEqual(row[4],"Home")
        net_sales=sum(row[13] or 0 for row in workbook["Orders"].iter_rows(min_row=5,values_only=True))
        self.assertAlmostEqual(net_sales,summary["kpis"]["sales"],places=2)
        chart_book=load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(chart_book["Categories"]._charts),1)
        workbook.close();chart_book.close()
        exported=self.client.get("/api/export/orders.csv?status=Cancelled&search=ORD")
        records=list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))
        self.assertTrue(records)
        self.assertTrue(all(row["status"]=="Cancelled" for row in records))
        payload=csv_bytes([{"product":"=HYPERLINK(\"bad\")"}],["product"]).decode("utf-8-sig")
        self.assertIn("'",payload)


if __name__=="__main__":
    unittest.main(verbosity=2)
