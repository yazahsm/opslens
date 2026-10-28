"""File parsing, explicit validation, and versioned local storage."""
import csv
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import uuid
import zipfile

ORDER_COLUMNS = ["order_id", "order_date", "sku", "product", "region", "category", "channel", "units",
                 "unit_price", "unit_cost", "status", "delivery_days", "target_delivery_days"]
INVENTORY_COLUMNS = ["sku", "product", "region", "category", "stock_on_hand", "reorder_point",
                     "avg_daily_units", "unit_cost", "snapshot_date"]
MAX_ROWS = 10000


def columns(kind):
    if kind not in ("orders", "inventory"):
        raise ValueError("Choose orders or inventory")
    return ORDER_COLUMNS if kind == "orders" else INVENTORY_COLUMNS


def parse_file(filename, payload, kind):
    if not payload or len(payload) > 5 * 1024 * 1024:
        raise ValueError("Choose a non-empty CSV or XLSX file smaller than 5 MB")
    if filename.lower().endswith(".csv"):
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Save the CSV with UTF-8 encoding")
        try:
            dialect = csv.Sniffer().sniff(text[:10000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        try:
            matrix = list(csv.reader(io.StringIO(text), dialect))
        except csv.Error:
            raise ValueError("The CSV is malformed or contains an excessively long cell")
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("The workbook expands beyond the 50 MB limit")
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
            try:
                sheet = book.worksheets[0]
                if sheet.max_column and sheet.max_column > 100:
                    raise ValueError("Import at most 100 columns")
                matrix = []
                for row in sheet.iter_rows(values_only=True):
                    matrix.append([value.isoformat()[:10] if isinstance(value, (date, datetime)) else value for value in row])
                    if len(matrix) > MAX_ROWS + 1:
                        raise ValueError("Import at most 10,000 records at once")
            finally:
                book.close()
        except (zipfile.BadZipFile, KeyError, IndexError):
            raise ValueError("This file is not a readable XLSX workbook")
        except Exception as exc:
            if isinstance(exc, ValueError):
                raise
            raise ValueError("The workbook could not be read. Use a plain XLSX file with values, not macros or formulas.") from exc
    else:
        raise ValueError("Supported formats: .csv and .xlsx")
    if not matrix:
        raise ValueError("The file has no header row")
    header = [str(value or "").strip().lower() for value in matrix[0]]
    while header and not header[-1]:
        header.pop()
    if len(header) != len(set(header)) or "" in header:
        raise ValueError("Column names must be non-empty and unique")
    missing = set(columns(kind)) - set(header)
    if missing:
        raise ValueError("Missing columns: " + ", ".join(sorted(missing)))
    rows = []
    for line, values in enumerate(matrix[1:], 2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(values) > len(header) and any(str(v or "").strip() for v in values[len(header):]):
            raise ValueError(f"Row {line} has more values than the header")
        rows.append({key: values[i] if i < len(values) and values[i] is not None else "" for i, key in enumerate(header)})
    if not rows:
        raise ValueError("The file has no data rows")
    if len(rows) > MAX_ROWS:
        raise ValueError("Import at most 10,000 records at once")
    return rows


def clean(rows, kind):
    """Keep the first valid key; exclude later duplicates and invalid rows."""
    required = columns(kind)
    accepted, rejected, repairs = [], [], []
    seen = set()
    for index, raw in enumerate(rows, 2):
        row, errors, changes = {}, [], []
        def text(key, title=False):
            value = str(raw.get(key, "") or "")
            fixed = value.strip()
            if title:
                fixed = fixed.title()
            if not fixed:
                errors.append(f"{key}: missing value")
            if len(fixed) > 200:
                errors.append(f"{key}: maximum 200 characters")
            if value != fixed:
                changes.append(key)
            row[key] = fixed
            return fixed

        def number(key, minimum=0, integer=False, optional=False):
            value = raw.get(key, "")
            if optional and (value is None or str(value).strip() == ""):
                row[key] = None
                return
            try:
                result = Decimal(str(value).strip())
                if not result.is_finite() or result < minimum or result > 1000000000:
                    raise ValueError()
                if integer and result != result.to_integral_value():
                    raise ValueError()
                if not integer and result != result.quantize(Decimal("0.01")):
                    raise ValueError()
                row[key] = int(result) if integer else float(result)
            except (InvalidOperation, ValueError, TypeError):
                errors.append(f"{key}: expected {'whole number' if integer else 'number (up to 2 decimals)'} >= {minimum}")
                row[key] = None

        date_key = "order_date" if kind == "orders" else "snapshot_date"
        for key in (["order_id", "sku", "product", "channel"] if kind == "orders" else ["sku", "product"]):
            text(key)
        for key in ("region", "category"):
            text(key, title=True)
        value = str(raw.get(date_key, "") or "").strip()
        try:
            if len(value) != 10:
                raise ValueError()
            row[date_key] = date.fromisoformat(value).isoformat()
            if date.fromisoformat(value) > date.today():
                raise ValueError()
        except ValueError:
            errors.append(f"{date_key}: expected a valid YYYY-MM-DD date, not in the future")
        number("unit_cost")
        if kind == "orders":
            text("status", title=True)
            if row["status"] not in ("Delivered", "Processing", "Cancelled"):
                errors.append("status: use Delivered, Processing or Cancelled")
            number("units", 1, integer=True)
            number("unit_price")
            number("target_delivery_days", 1, integer=True)
            number("delivery_days", 0, integer=True, optional=row["status"] != "Delivered")
            key = row["order_id"]
        else:
            number("stock_on_hand", integer=True)
            number("reorder_point", integer=True)
            number("avg_daily_units")
            key = (row["sku"], row["region"])
        if key in seen:
            errors.append("duplicate key: first valid record retained")
        if errors:
            rejected.append({"row": index, "key": str(raw.get("order_id" if kind == "orders" else "sku", "")),
                             "reasons": errors, "raw": {k: raw.get(k, "") for k in required}})
            continue
        seen.add(key)
        accepted.append(row)
        if changes:
            repairs.append({"row": index, "fields": changes})
    return {"rows": accepted, "rejected": rejected, "repairs": repairs,
            "raw_count": len(rows), "accepted_count": len(accepted), "rejected_count": len(rejected),
            "normalized_count": len(repairs), "pass_rate": round(len(accepted) / len(rows) * 100, 2) if rows else 0}


class Store:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript("""CREATE TABLE IF NOT EXISTS datasets (
                id TEXT PRIMARY KEY, kind TEXT NOT NULL, name TEXT NOT NULL, created_at TEXT NOT NULL,
                payload TEXT NOT NULL, source_hash TEXT NOT NULL, is_demo INTEGER NOT NULL);
                CREATE TABLE IF NOT EXISTS active (kind TEXT PRIMARY KEY, dataset_id TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS previews (id TEXT PRIMARY KEY, created_at TEXT, payload TEXT);""")

    @contextmanager
    def connect(self):
        connection = sqlite3.connect(self.path, timeout=20)
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def save(self, kind, name, result, is_demo=False, activate=True):
        identifier = uuid.uuid4().hex
        payload = json.dumps(result, allow_nan=False, default=str)
        created = datetime.now().astimezone().isoformat(timespec="seconds")
        with self.connect() as db:
            db.execute("INSERT INTO datasets VALUES (?,?,?,?,?,?,?)", (identifier, kind, name, created, payload,
                       hashlib.sha256(payload.encode()).hexdigest(), int(is_demo)))
            if activate:
                db.execute("INSERT OR REPLACE INTO active VALUES (?,?)", (kind, identifier))
        return identifier

    def get(self, kind):
        columns(kind)
        with self.connect() as db:
            record = db.execute("SELECT d.id,d.name,d.created_at,d.payload,d.source_hash,d.is_demo FROM datasets d JOIN active a ON d.id=a.dataset_id WHERE a.kind=?", (kind,)).fetchone()
        if record is None:
            return None
        identifier, name, created, payload, digest, demo = record
        result = json.loads(payload)
        result.update(id=identifier, kind=kind, name=name, created_at=created, source_hash=digest, is_demo=bool(demo))
        return result

    def preview(self, kind, name, result):
        token = uuid.uuid4().hex
        with self.connect() as db:
            db.execute("DELETE FROM previews WHERE created_at < datetime('now', '-1 day')")
            db.execute("INSERT INTO previews VALUES (?,datetime('now'),?)", (token, json.dumps({"kind": kind, "name": name, "result": result}, default=str)))
        return token

    def apply(self, token):
        with self.connect() as db:
            found = db.execute("SELECT payload FROM previews WHERE id=? AND created_at >= datetime('now','-1 day')", (token,)).fetchone()
        if not found:
            raise ValueError("Import preview expired or was already applied. Preview the file again.")
        payload = json.loads(found[0])
        if not payload["result"]["accepted_count"]:
            raise ValueError("No valid records to import")
        identifier = self.save(payload["kind"], payload["name"], payload["result"])
        with self.connect() as db:
            db.execute("DELETE FROM previews WHERE id=?", (token,))
        return identifier


def metadata(dataset):
    return {key: dataset[key] for key in ("id", "kind", "name", "created_at", "source_hash", "is_demo", "raw_count",
                                         "accepted_count", "rejected_count", "normalized_count", "pass_rate")}
